import pandas as pd
import openpyxl
import os
import sys

sys.stdout.reconfigure(encoding='utf-8')

INPUT = r'D:\Tool-Glm\file'
OUT = r'D:\Tool-Glm\HashtagNew\Master_Hashtag_DB.csv'

FILES = {
    'W1': 'W1.xlsx', '3D1': '3D1.xlsx', '3D2': '3D2.xlsx',
    '3D3': '3D3.xlsx', 'Zipper': 'Zipper.xlsx', 'Charging': 'Charging.xlsx',
}
APPS = list(FILES.keys())

master = {}

def add(h, p, cat, app):
    h = str(h).strip().lower().replace(' ', '')
    p = str(p).strip().lower().replace(' ', '') if pd.notna(p) and str(p).strip() else None
    if not h or h in ('nan', 'none', ''):
        return
    if h not in master:
        master[h] = {'parent': p, 'category': cat, 'apps': set([app])}
    else:
        master[h]['apps'].add(app)
        if master[h]['parent'] is None and p is not None:
            master[h]['parent'] = p

for app, fn in FILES.items():
    fp = os.path.join(INPUT, fn)
    print(f"Processing {app}...")
    wb = openpyxl.load_workbook(fp, read_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))

    ds = 0
    for i, r in enumerate(rows):
        c0 = str(r[0]).strip().lower() if r[0] else ''
        if c0 in ('sport', 'object'):
            ds = i
            break

    pl1 = pl2 = None
    for row in rows[ds:]:
        v = [str(x).strip() if x is not None else '' for x in row]
        if len(v) < 9:
            continue

        l1 = v[0] if v[0].lower() not in ('', 'nan', 'none', 'object', 'style', 'color') else ''
        l2 = v[1] if v[1].lower() not in ('', 'nan', 'none') else ''
        l3 = v[2] if v[2].lower() not in ('', 'nan', 'none') else ''
        sv = v[5] if v[5].lower() not in ('', 'nan', 'none', 'style') else ''
        cv = v[8] if v[8].lower() not in ('', 'nan', 'none', 'color') else ''

        if not any([l1, l2, l3, sv, cv]):
            continue

        if l1:
            pl1 = l1
            pl2 = None
            add(l1, None, 'object', app)
        else:
            l1 = pl1

        if l2:
            pl2 = l2
            add(l2, l1, 'object', app)
        else:
            l2 = pl2

        if l3:
            add(l3, l2, 'object', app)
        if sv:
            add(sv, None, 'style', app)
        if cv:
            add(cv, None, 'color', app)
    wb.close()

rows_out = []
COL_MAP = {'W1': 'w1', '3D1': 'd3d1', '3D2': 'd3d2', '3D3': 'd3d3', 'Zipper': 'zipper', 'Charging': 'charging'}
for h, info in master.items():
    row = {'hashtag': h, 'parent_hashtag': info['parent'] or '', 'category': info['category']}
    for a in APPS:
        row[COL_MAP[a]] = a in info['apps']
    rows_out.append(row)

df = pd.DataFrame(rows_out)
df = df[~df['hashtag'].str.contains('unnamed', na=False)]
df = df[~df['hashtag'].str.match(r'^\*+$', na=False)]

cat_ord = {'object': 1, 'style': 2, 'color': 3}
df['_s'] = df['category'].map(cat_ord)
df = df.sort_values(['_s', 'hashtag']).drop(columns=['_s'])

cols = ['hashtag', 'parent_hashtag', 'category', 'w1', 'd3d1', 'd3d2', 'd3d3', 'zipper', 'charging']
df = df[cols]

df.to_csv(OUT, index=False, encoding='utf-8-sig')

n_obj = len(df[df['category'] == 'object'])
n_sty = len(df[df['category'] == 'style'])
n_clr = len(df[df['category'] == 'color'])
print(f'Done: {len(df)} rows -> {OUT}')
print(f'Object: {n_obj}, Style: {n_sty}, Color: {n_clr}')
